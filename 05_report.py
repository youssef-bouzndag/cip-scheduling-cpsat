import pandas as pd
from openpyxl.styles import Font, PatternFill

SETUP_MIN = 15  # must match 03_scheduler.py

schedule = pd.read_csv("schedule.csv", parse_dates=["start", "end"])
production = pd.read_csv(
    "production.csv",
    parse_dates=["earliest_clean_start", "latest_clean_start"]
)
options = pd.read_csv("cleaning_tasks.csv")[["batch_id", "cleaning_resource"]]

# ---------- 1. One row per batch ----------
batches = schedule.merge(
    production[["batch_id", "equipment", "earliest_clean_start",
                "latest_clean_start", "slack_min"]],
    on="batch_id"
)
batches = batches.rename(columns={"earliest_clean_start": "ready",
                                  "latest_clean_start": "latest_start"})
batches["cleaning_min"] = (batches["end"] - batches["start"]).dt.total_seconds() / 60
batches["slack_left_min"] = batches["slack_min"] - batches["waiting_min"]
batches["status"] = batches["waiting_min"].apply(lambda w: "WAITS" if w > 0 else "on time")

# Other resources that could have cleaned this batch
def other_resources(batch, resource):
    possible = options[options["batch_id"] == batch]["cleaning_resource"].tolist()
    others = [x for x in possible if x != resource]
    return ", ".join(others) if others else "none (only option)"

batches["other_resources"] = [
    other_resources(b, r) for b, r in zip(batches["batch_id"], batches["resource"])
]

batches = batches.sort_values(["resource", "start"]).reset_index(drop=True)

# Schedule span: first equipment ready -> last resource free again
span_start = batches["ready"].min()
span_end = batches["end"].max() + pd.Timedelta(minutes=SETUP_MIN)
span_min = (span_end - span_start).total_seconds() / 60

# ---------- 2. Why does a batch wait? + order on each resource ----------
batches["blocked_by"] = ""
batches["reason"] = ""
timeline = []

for r in batches["resource"].unique():
    sub = batches[batches["resource"] == r].sort_values("start")
    prev = None
    order = 0
    for i, row in sub.iterrows():
        order += 1
        if prev is None:
            idle_before = (row["start"] - span_start).total_seconds() / 60
        else:
            prev_free = prev["end"] + pd.Timedelta(minutes=SETUP_MIN)
            idle_before = (row["start"] - prev_free).total_seconds() / 60
            if row["waiting_min"] > 0 and prev_free > row["ready"]:
                batches.loc[i, "blocked_by"] = prev["batch_id"]
                batches.loc[i, "reason"] = (
                    f"{r} busy with {prev['batch_id']} until {prev_free:%H:%M} "
                    f"(cleaning ends {prev['end']:%H:%M} + {SETUP_MIN} min setup)"
                )
        timeline.append({
            "resource": r, "order": order, "batch_id": row["batch_id"],
            "equipment": row["equipment"], "start": row["start"], "end": row["end"],
            "free_again_at": row["end"] + pd.Timedelta(minutes=SETUP_MIN),
            "idle_before_min": idle_before, "waiting_min": row["waiting_min"],
        })
        prev = row

# A waiting batch with no blocker found means something is wrong
no_reason = (batches["waiting_min"] > 0) & (batches["reason"] == "")
batches.loc[no_reason, "reason"] = "no blocker found, check the model"

timeline = pd.DataFrame(timeline)

# ---------- 3. Sheet: batches that wait ----------
waiting = batches[batches["waiting_min"] > 0][
    ["batch_id", "equipment", "resource", "ready", "start", "waiting_min",
     "latest_start", "slack_left_min", "blocked_by", "reason", "other_resources"]
].sort_values("waiting_min", ascending=False)

if len(waiting) == 0:
    waiting = pd.DataFrame({"message": ["No batch waits."]})

# ---------- 4. Sheet: workload per cleaning resource ----------
workload = batches.groupby("resource").agg(
    cleanings=("batch_id", "count"),
    cleaning_min=("cleaning_min", "sum"),
    batches_waiting=("waiting_min", lambda s: int((s > 0).sum())),
    waiting_min=("waiting_min", "sum"),
    first_start=("start", "min"),
    last_end=("end", "max"),
).reset_index()
workload["setup_min"] = workload["cleanings"] * SETUP_MIN
workload["busy_min"] = workload["cleaning_min"] + workload["setup_min"]
workload["idle_min"] = span_min - workload["busy_min"]
workload["busy_%"] = (workload["busy_min"] / span_min * 100).round(1)
workload["idle_%"] = (workload["idle_min"] / span_min * 100).round(1)
workload["order"] = workload["resource"].apply(
    lambda r: " -> ".join(batches[batches["resource"] == r].sort_values("start")["batch_id"])
)

# ---------- 5. Sheet: workload per equipment ----------
equipment = batches.groupby("equipment").agg(
    cleanings=("batch_id", "count"),
    cleaning_min=("cleaning_min", "sum"),
    waiting_min=("waiting_min", "sum"),
    tightest_slack_left_min=("slack_left_min", "min"),
).reset_index()

# ---------- 6. Sheet: summary ----------
busiest = workload.sort_values("busy_%", ascending=False).iloc[0]
tightest = batches.sort_values("slack_left_min").iloc[0]
summary = pd.DataFrame([
    {"metric": "Total cleanings", "value": len(batches)},
    {"metric": "Total waiting time (min)", "value": int(batches["waiting_min"].sum())},
    {"metric": "Batches that wait", "value": int((batches["waiting_min"] > 0).sum())},
    {"metric": "Which batches wait", "value": ", ".join(
        batches[batches["waiting_min"] > 0]["batch_id"]) or "none"},
    {"metric": "Longest wait (min)", "value": int(batches["waiting_min"].max())},
    {"metric": "Busy time per resource (min)", "value": ", ".join(
        f"{r}: {int(m)}" for r, m in zip(workload["resource"], workload["busy_min"]))},
    {"metric": "Idle time per resource (min)", "value": ", ".join(
        f"{r}: {int(m)}" for r, m in zip(workload["resource"], workload["idle_min"]))},
    {"metric": "Busiest resource", "value": f"{busiest['resource']} ({busiest['busy_%']}%)"},
    {"metric": "Tightest batch (least slack left)",
     "value": f"{tightest['batch_id']} ({int(tightest['slack_left_min'])} min left)"},
    {"metric": "Setup time per cleaning (min)", "value": SETUP_MIN},
    {"metric": "Schedule span (min)", "value": int(span_min)},
])

# ---------- 7. Sheet: definitions ----------
definitions = pd.DataFrame([
    {"term": "Waiting", "about": "equipment",
     "meaning": "Minutes between equipment being ready (production ended) and its cleaning starting"},
    {"term": "Slack", "about": "batch",
     "meaning": "Size of the allowed start window: latest start - earliest start (from the input data)"},
    {"term": "Slack left", "about": "batch",
     "meaning": "Slack - waiting. Room still left after scheduling; small = fragile"},
    {"term": "Busy", "about": "cleaning resource",
     "meaning": "Cleaning time + setup time. The resource cannot take another job"},
    {"term": "Idle", "about": "cleaning resource",
     "meaning": "Schedule span - busy time. Spare capacity, but possibly at the wrong hours"},
    {"term": "Span", "about": "whole schedule",
     "meaning": "From the first equipment ready time to the last time a resource is free again"},
    {"term": "Blocked by", "about": "batch",
     "meaning": "The previous cleaning on the same resource that is still running or flushing"},
])

# ---------- 8. Write Excel ----------
all_batches = batches[["batch_id", "equipment", "resource", "status", "ready", "start",
                       "end", "latest_start", "cleaning_min", "waiting_min",
                       "slack_min", "slack_left_min", "blocked_by", "reason",
                       "other_resources"]]

sheets = {
    "Summary": summary,
    "Waiting batches": waiting,
    "Resource workload": workload,
    "Resource timeline": timeline,
    "Equipment": equipment,
    "All batches": all_batches,
    "Definitions": definitions,
}

amber = PatternFill("solid", fgColor="FFE699")
with pd.ExcelWriter("cleaning_report.xlsx") as writer:
    for name, df in sheets.items():
        df.to_excel(writer, sheet_name=name, index=False)
        ws = writer.sheets[name]
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for i, col in enumerate(df.columns):
            width = max(len(str(col)), df[col].astype(str).str.len().max()) + 2
            ws.column_dimensions[chr(65 + i)].width = min(width, 80)
        # Highlight rows with waiting time
        if "waiting_min" in df.columns:
            wi = list(df.columns).index("waiting_min")
            for row in ws.iter_rows(min_row=2):
                if (row[wi].value or 0) > 0:
                    for cell in row:
                        cell.fill = amber

# ---------- 9. Print the key parts ----------
print(summary.to_string(index=False))
print("\nBatches that wait:")
print(waiting.to_string(index=False))
print("\nWorkload per resource:")
print(workload[["resource", "cleanings", "cleaning_min", "setup_min", "busy_min",
                "idle_min", "busy_%", "idle_%", "order"]].to_string(index=False))
print("\nSaved cleaning_report.xlsx")